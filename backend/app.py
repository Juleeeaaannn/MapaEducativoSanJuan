"""
API de establecimientos educativos de San Juan.

Lee el archivo datos_actualizados.xlsx (raíz del proyecto) con openpyxl,
detecta automáticamente las columnas (CUEANEXO, latitud, longitud, nombre, etc.)
y expone los datos para el mapa interactivo.

Endpoints:
    GET /api/health              -> estado de la API
    GET /api/meta                -> info del archivo y columnas detectadas
    GET /api/establecimientos    -> todos los registros para el mapa
    GET /                        -> frontend compilado (scripts/3_produccion.bat)

Ejecución:  python -m uvicorn backend.app:app --port 8000
"""
from __future__ import annotations

import os
import re
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = Path(os.environ.get("DATOS_XLSX", str(ROOT / "datos_actualizados.xlsx")))
DIST_DIR = ROOT / "frontend" / "dist"

# --------------------------------------------------------------------------
# Detección automática de columnas
# --------------------------------------------------------------------------

def _norm(value: Any) -> str:
    """minúsculas, sin acentos y sin símbolos, para comparar encabezados."""
    if value is None:
        return ""
    s = unicodedata.normalize("NFKD", str(value))
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^a-z0-9]+", "_", s.strip().lower())
    return s.strip("_")


ALIASES = {
    "cueanexo": ("cueanexo", "cueanexos", "cue_anexo", "cue_anexos", "id_cueanexo", "cueanexo_id"),
    "cue": ("cue",),
    "nombre": (
        "nombre", "nombres", "nombre_establecimiento", "establecimiento",
        "nombre_institucion", "institucion", "denominacion", "escuela",
        "nombre_escuela", "nombre_del_establecimiento",
    ),
    "latitud": ("latitud", "lat", "latitud_decimal", "coordenada_latitud", "latitud_y", "lat_y", "y"),
    "longitud": (
        "longitud", "long", "lon", "lng", "longitud_decimal",
        "coordenada_longitud", "longitud_x", "long_x", "x",
    ),
    "coordenadas": (
        "coordenadas", "coordenada", "coords", "coord",
        "latitud_longitud", "lat_lon", "latlng", "lat_long",
    ),
    "departamento": ("departamento", "depto", "dpto", "departamento_educativo"),
    "localidad": ("localidad", "ciudad"),
    "nivel": ("nivel", "niveles", "nivel_ensenanza", "nivel_servicio", "nivel_oferta", "ofertas"),
    "sector": ("sector", "sector_gestion", "gestion", "ambito", "ambito_gestion"),
    "direccion": ("direccion", "domicilio", "calle", "direccion_postal"),
}


def _match_alias(header_norm: str, aliases: tuple[str, ...]) -> bool:
    if not header_norm:
        return False
    for a in aliases:
        if header_norm == a or header_norm.startswith(a + "_") or header_norm.endswith("_" + a):
            return True
    return False


def _es_alias_conocido(header_norm: str) -> bool:
    return any(_match_alias(header_norm, al) for al in ALIASES.values())


# --------------------------------------------------------------------------
# Lectura del Excel
# --------------------------------------------------------------------------

def _limpiar_texto(value: Any, max_len: int = 160) -> str:
    if value is None:
        return ""
    return str(value).strip()[:max_len]


def _a_float(value: Any) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace("\u00a0", "").replace(" ", "")
    if not s:
        return None
    s = s.replace(",", ".")  # coma decimal -> punto
    try:
        return float(s)
    except ValueError:
        return None


def _a_cue(value: Any) -> str:
    """CUEANEXO solo dígitos (soporta celdas numéricas tipo 700012500.0)."""
    if value is None:
        return ""
    s = str(value).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return re.sub(r"\D", "", s)


def _parsear_coordenadas(value: Any) -> tuple[Optional[float], Optional[float]]:
    """Parsea una celda con latitud y longitud juntas: 'lat, lon', 'lat;lon' o 'lat lon'."""
    if value is None:
        return None, None
    s = str(value).strip().replace("\u00a0", "")
    if not s:
        return None, None

    def _dos(par: list[str]) -> tuple[Optional[float], Optional[float]] | tuple[None, None]:
        if len(par) != 2:
            return None, None
        lat, lon = _a_float(par[0]), _a_float(par[1])
        if lat is None or lon is None:
            return None, None
        return lat, lon

    for sep in (";", ","):
        if sep in s:
            lat, lon = _dos([p.strip() for p in s.split(sep) if p.strip()])
            if lat is not None and lon is not None:
                return lat, lon

    lat, lon = _dos(s.split())
    if lat is not None and lon is not None:
        return lat, lon
    return None, None


def _detectar_encabezados(filas: list[tuple], max_filas: int = 20):
    """Busca en las primeras filas la que más encabezados conocidos tiene."""
    mejor_score, mejor_idx = 0, 0
    for idx, fila in enumerate(filas[:max_filas]):
        score = sum(1 for c in fila if _es_alias_conocido(_norm(c)))
        if score > mejor_score:
            mejor_score, mejor_idx = score, idx
    return mejor_score, mejor_idx


def _mapear_columnas(fila_encabezado):
    """Devuelve {campo: índice de columna} y la lista de columnas extra."""
    mapeo: dict[str, int] = {}
    usadas: set[int] = set()
    for campo, aliases in ALIASES.items():
        for i, celda in enumerate(fila_encabezado):
            if i in usadas:
                continue
            if _match_alias(_norm(celda), aliases):
                mapeo[campo] = i
                usadas.add(i)
                break
    extras = []
    for i, celda in enumerate(fila_encabezado):
        if i not in usadas and celda not in (None, ""):
            extras.append((i, _limpiar_texto(celda, 60) or f"col_{i + 1}"))
    return mapeo, extras


def cargar_datos() -> dict:
    """Lee el Excel y devuelve los establecimientos + estadísticas."""
    if not XLSX_PATH.exists():
        return {
            "ok": False,
            "error": f"No se encontró '{XLSX_PATH.name}' en la raíz del proyecto.",
            "archivo": str(XLSX_PATH),
        }

    try:
        wb = load_workbook(XLSX_PATH, read_only=True, data_only=True)
    except Exception as exc:  # pragma: no cover
        return {"ok": False, "error": f"No se pudo abrir el Excel: {exc}", "archivo": str(XLSX_PATH)}

    # Leer cada hoja una sola vez y elegir la de mejor encabezado.
    # (En modo read_only conviene evitar reiterar la misma hoja.)
    MAX_FILAS = 200_000
    mejor_hoja, mejor_score = None, -1
    mejor_filas: list[tuple] = []
    mejor_idx = 0

    for nombre_hoja in wb.sheetnames:
        ws = wb[nombre_hoja]
        filas = [
            tuple(f)
            for f in ws.iter_rows(min_row=1, max_row=MAX_FILAS, values_only=True)
        ]
        score, idx = _detectar_encabezados(filas)
        if score > mejor_score:
            mejor_hoja, mejor_score = nombre_hoja, score
            mejor_filas, mejor_idx = filas, idx

    if mejor_hoja is None or mejor_score < 1:
        return {
            "ok": False,
            "error": (
                "No se detectaron columnas conocidas en el Excel. "
                "Revisá que exista una fila con encabezados (CUEANEXO, LATITUD, LONGITUD...). "
                "Corré scripts/4_verificar_datos.bat para ver el detalle."
            ),
            "archivo": str(XLSX_PATH),
        }

    filas = mejor_filas
    idx_enc = mejor_idx
    fila_enc = filas[idx_enc]
    mapeo, extras_cols = _mapear_columnas(fila_enc)

    if "cueanexo" not in mapeo and "nombre" not in mapeo:
        encabezados = [_limpiar_texto(c, 40) for c in fila_enc if c not in (None, "")]
        return {
            "ok": False,
            "error": (
                "No se encontró una columna de CUEANEXO ni de NOMBRE. "
                f"Encabezados detectados en la fila {idx_enc + 1}: {encabezados}. "
                "Podés ajustar los alias en backend/app.py (diccionario ALIASES)."
            ),
            "archivo": str(XLSX_PATH),
        }

    establecimientos = []
    ids_vistos: dict[str, int] = {}
    n_datos = 0
    invalidas = 0

    for fila in filas[idx_enc + 1 :]:
        if fila is None or all(
            c is None or (isinstance(c, str) and not c.strip()) for c in fila
        ):
            continue
        n_datos += 1

        def celda(campo: str) -> Any:
            i = mapeo.get(campo)
            if i is None or i >= len(fila):
                return None
            return fila[i]

        lat = _a_float(celda("latitud"))
        lon = _a_float(celda("longitud"))
        if lat is None or lon is None:
            # Fallback: latitud y longitud juntas en una sola celda ('lat, lon').
            lat_c, lon_c = _parsear_coordenadas(celda("coordenadas"))
            if lat is None:
                lat = lat_c
            if lon is None:
                lon = lon_c
        coords_ok = (
            lat is not None
            and lon is not None
            and -90 <= lat <= 90
            and -180 <= lon <= 180
        )
        if (lat is not None or lon is not None) and not coords_ok:
            invalidas += 1

        cueanexo = _a_cue(celda("cueanexo"))
        nombre = _limpiar_texto(celda("nombre"))
        base_id = cueanexo or nombre or f"fila-{n_datos}"
        n_rep = ids_vistos.get(base_id, 0)
        ids_vistos[base_id] = n_rep + 1
        est_id = base_id if n_rep == 0 else f"{base_id}-{n_rep}"

        extras = {}
        for i, nombre_col in extras_cols:
            if i < len(fila):
                valor = _limpiar_texto(fila[i])
                if valor:
                    extras[nombre_col] = valor

        establecimientos.append(
            {
                "id": est_id,
                "cueanexo": cueanexo,
                "cue": cueanexo[:7] if len(cueanexo) >= 7 else "",
                "nombre": nombre,
                "lat": lat if coords_ok else None,
                "lon": lon if coords_ok else None,
                "coordenadas_validas": coords_ok,
                "departamento": _limpiar_texto(celda("departamento")),
                "localidad": _limpiar_texto(celda("localidad")),
                "nivel": _limpiar_texto(celda("nivel")),
                "sector": _limpiar_texto(celda("sector")),
                "direccion": _limpiar_texto(celda("direccion")),
                "extras": extras,
            }
        )

    con_coords = sum(1 for e in establecimientos if e["coordenadas_validas"])
    try:
        mtime = XLSX_PATH.stat().st_mtime
        actualizado = datetime.fromtimestamp(mtime).isoformat(timespec="seconds")
    except OSError:
        actualizado = None

    return {
        "ok": True,
        "archivo": XLSX_PATH.name,
        "actualizado": actualizado,
        "hoja": mejor_hoja,
        "fila_encabezados": idx_enc + 1,
        "columnas": {
            campo: _limpiar_texto(fila_enc[i], 60) for campo, i in mapeo.items()
        },
        "columnas_extra": [nombre for _, nombre in extras_cols],
        "total_filas": n_datos,
        "con_coordenadas": con_coords,
        "coordenadas_invalidas": invalidas,
        "sin_coordenadas": n_datos - con_coords,
        "cueanexos_duplicados": sum(1 for v in ids_vistos.values() if v > 1),
        "establecimientos": establecimientos,
    }


_cache: dict[str, Any] = {"mtime": None, "data": None}


def _datos_cached() -> dict:
    """Relee el Excel solo si cambió (mtime)."""
    try:
        mtime: Optional[float] = XLSX_PATH.stat().st_mtime
    except OSError:
        mtime = None
    if _cache["data"] is None or _cache["mtime"] != mtime:
        _cache["data"] = cargar_datos()
        _cache["mtime"] = mtime
    return _cache["data"]


# --------------------------------------------------------------------------
# API
# --------------------------------------------------------------------------

app = FastAPI(title="Establecimientos Educativos · San Juan", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.get("/api/health")
def health():
    return {"ok": True, "servicio": "escuelas-san-juan-api"}


@app.get("/api/meta")
def meta():
    datos = _datos_cached()
    return {k: v for k, v in datos.items() if k != "establecimientos"}


# --------------------------------------------------------------------------
# Entrecruzamientos: subida de archivo (cueanexo, entrecruzamiento)
# --------------------------------------------------------------------------

MAX_ENTRECRUZAMIENTOS = 5_000


def _leer_pares_cue(archivo: UploadFile) -> list[tuple[str, str]]:
    """Lee xlsx o csv y devuelve pares (cueanexo, entrecruzamiento) con solo dígitos."""
    contenido = archivo.file.read()
    nombre = (archivo.filename or "").lower()

    filas: list[list[Any]]
    if nombre.endswith((".xlsx", ".xlsm")):
        import io

        try:
            wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        except Exception as exc:
            raise HTTPException(400, f"No se pudo leer el Excel: {exc}")
        ws = wb[wb.sheetnames[0]]
        filas = [list(f) for f in ws.iter_rows(max_row=MAX_ENTRECRUZAMIENTOS + 50, values_only=True)]
    elif nombre.endswith(".csv"):
        import csv
        import io

        texto = contenido.decode("utf-8-sig", errors="replace")
        sniffer = csv.Sniffer()
        try:
            dialect = sniffer.sniff(texto[:2048])
            delimitador = dialect.delimiter
        except csv.Error:
            delimitador = ";" if texto.count(";") > texto.count(",") else ","
        filas = list(csv.reader(io.StringIO(texto), delimiter=delimitador))
    else:
        raise HTTPException(400, "Formato no soportado. Usá .xlsx o .csv")

    # Detectar la fila de encabezados: la que contenga cueanexo (y entrecruzamiento).
    inicio = 0
    for i, fila in enumerate(filas[:20]):
        normalizadas = {_norm(c) for c in fila if c is not None}
        if any(h.startswith("cueanexo") for h in normalizadas):
            inicio = i
            break

    col_cue, col_ent = None, None
    for i, celda in enumerate(filas[inicio]):
        h = _norm(celda)
        if h.startswith("cueanexo") and col_cue is None:
            col_cue = i
        elif h.startswith("entrecruzamiento") and col_ent is None:
            col_ent = i
    if col_cue is None or col_ent is None:
        raise HTTPException(
            400,
            "No se encontraron las columnas 'cueanexo' y 'entrecruzamiento' en el archivo.",
        )

    pares: list[tuple[str, str]] = []
    for fila in filas[inicio + 1 :]:
        if not fila or all(c is None or str(c).strip() == "" for c in fila):
            continue
        cue = _a_cue(fila[col_cue]) if col_cue < len(fila) else ""
        ent = _a_cue(fila[col_ent]) if col_ent < len(fila) else ""
        if cue and ent:
            pares.append((cue, ent))
    return pares


@app.post("/api/entrecruzamientos")
async def subir_entrecruzamientos(archivo: UploadFile = File(...)):
    """Recibe un xlsx/csv (cueanexo, entrecruzamiento) y devuelve las líneas a dibujar.

    Los valores pueden ser CUEANEXO completo (p. ej. 9 dígitos) o CUE base
    de 7 dígitos (solo si corresponde a una única escuela, para evitar ambigüedades).
    """
    pares = _leer_pares_cue(archivo)
    if not pares:
        raise HTTPException(400, "El archivo no contiene filas válidas (cueanexo y entrecruzamiento).")

    datos = _datos_cached()
    if not datos.get("ok"):
        raise HTTPException(500, "El Excel principal no está disponible.")

    indice_cueanexo: dict[str, dict] = {}
    indice_cue7: dict[str, list[dict]] = {}
    for e in datos["establecimientos"]:
        if e["coordenadas_validas"] and len(e["cueanexo"]) >= 7:
            indice_cueanexo[e["cueanexo"]] = e
            indice_cue7.setdefault(e["cueanexo"][:7], []).append(e)

    def _buscar(cue: str) -> Optional[dict]:
        """CUEANEXO exacto, o CUE de 7 dígitos solo si es inequívoco."""
        if cue in indice_cueanexo:
            return indice_cueanexo[cue]
        if len(cue) == 7:
            candidatos = indice_cue7.get(cue, [])
            if len(candidatos) == 1:
                return candidatos[0]
        return None

    lineas, ignorados = [], {"cue_no_encontrado": set(), "formato_invalido": 0}
    for cue, ent in pares:
        if len(cue) < 7 or len(ent) < 7:
            ignorados["formato_invalido"] += 1
            continue
        origen = _buscar(cue)
        destino = _buscar(ent)
        if not origen or not destino:
            ignorados["cue_no_encontrado"].add(cue if not origen else ent)
            continue
        lineas.append(
            {
                "id": f"{cue}-{ent}",
                "origen": cue,
                "destino": ent,
                "coordenadas": [
                    [origen["lat"], origen["lon"]],
                    [destino["lat"], destino["lon"]],
                ],
            }
        )
        if len(lineas) >= MAX_ENTRECRUZAMIENTOS:
            break

    return {
        "ok": True,
        "total_filas": len(pares),
        "lineas": lineas,
        "ignorados": {
            "formato_invalido": ignorados["formato_invalido"],
            "cue_no_encontrado": sorted(ignorados["cue_no_encontrado"])[:50],
            "cue_no_encontrado_total": len(ignorados["cue_no_encontrado"]),
        },
    }


@app.get("/api/establecimientos")
def establecimientos():
    datos = _datos_cached()
    if not datos.get("ok"):
        return JSONResponse(status_code=500, content=datos)
    return datos


@app.get("/", include_in_schema=False)
def raiz():
    index = DIST_DIR / "index.html"
    if index.exists():
        return FileResponse(index)
    return JSONResponse(
        {
            "ok": True,
            "mensaje": (
                "API activa. Levantá el frontend con scripts/2_iniciar.bat "
                "(desarrollo) o scripts/3_produccion.bat (servidor único)."
            ),
        }
    )


# Si ya existe el build de producción, servirlo (API + frontend en un solo puerto)
if (DIST_DIR / "assets").exists():
    app.mount("/assets", StaticFiles(directory=DIST_DIR / "assets"), name="assets")
